import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme Colors
    BG_COLOR = RGBColor(255, 255, 255)       # White Background
    SAGE_GREEN = RGBColor(140, 156, 148)     # Sage Green bar (#8c9c94)
    TEXT_RED = RGBColor(219, 68, 54)         # Orange-Red Headings (#db4436)
    TEXT_DARK = RGBColor(30, 41, 59)         # Dark slate for primary text
    TEXT_MUTED = RGBColor(100, 116, 139)     # Grey for description paragraph

    blank_layout = prs.slide_layouts[6] # Blank layout

    # Helper function to set slide white background
    def set_white_bg(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background() # No border
        return bg

    # Helper function to add sage green top band
    def add_sage_bar(slide):
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.4)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAGE_GREEN
        bar.line.fill.background()

    # Helper function to create standard header
    def add_slide_header(slide, title_text):
        add_sage_bar(slide)
        
        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_RED

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (University Header, Department, Details)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide1)

    # Insert University Logo Header cropped from PDF Page 1
    logo_path = "scratch_extracted/university_logo_header.png"
    if os.path.exists(logo_path):
        # Spans width of the slide: 13.333 inches, aspect ratio is 960/140, height = 13.333 / 6.857 = 1.944
        slide1.shapes.add_picture(logo_path, Inches(0), Inches(0), Inches(13.333), Inches(1.944))
    else:
        # Fallback text block if logo header missing
        fallback_logo = slide1.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11.7), Inches(1.2))
        tf_fb = fallback_logo.text_frame
        p_fb = tf_fb.paragraphs[0]
        p_fb.text = "MALLA REDDY UNIVERSITY"
        p_fb.alignment = PP_ALIGN.CENTER
        p_fb.font.size = Pt(32)
        p_fb.font.bold = True
        p_fb.font.color.rgb = RGBColor(30, 60, 150)

    # Department Name
    dept_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(0.6))
    tf_dept = dept_box.text_frame
    tf_dept.word_wrap = True
    p_dept = tf_dept.paragraphs[0]
    p_dept.text = "Department of Computer Science & Engineering"
    p_dept.alignment = PP_ALIGN.CENTER
    p_dept.font.name = "Arial"
    p_dept.font.size = Pt(24)
    p_dept.font.bold = True
    p_dept.font.color.rgb = TEXT_RED

    # Project Title
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.0))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = "AthleteHub – Athlete Portal & Sports Management System"
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial"
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK

    # Bottom Right Block (Batch & Students)
    students_box = slide1.shapes.add_textbox(Inches(7.8), Inches(4.2), Inches(4.7), Inches(2.6))
    tf_stud = students_box.text_frame
    tf_stud.word_wrap = True
    
    p_batch = tf_stud.paragraphs[0]
    p_batch.text = "Batch No. 19"
    p_batch.font.name = "Arial"
    p_batch.font.size = Pt(17)
    p_batch.font.bold = True
    p_batch.font.color.rgb = TEXT_DARK
    p_batch.space_after = Pt(10)

    students = [
        "1. 2511CS010126 - M . Sai Ram",
        "2. 2511CS010092 - D . Sumanth",
        "3. 2511CS010136 - K . Divija",
        "4. 2511CS010146 - P.Raghavendra"
    ]
    for stud in students:
        p = tf_stud.add_paragraph()
        p.text = stud
        p.font.name = "Arial"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    # Bottom Left Block (Guide)
    guide_box = slide1.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(4.5), Inches(1.5))
    tf_guide = guide_box.text_frame
    tf_guide.word_wrap = True
    
    p_g_label = tf_guide.paragraphs[0]
    p_g_label.text = "Guide"
    p_g_label.font.name = "Arial"
    p_g_label.font.size = Pt(17)
    p_g_label.font.bold = True
    p_g_label.font.color.rgb = TEXT_DARK
    p_g_label.space_after = Pt(6)

    p_g_name = tf_guide.add_paragraph()
    p_g_name.text = "Mrs. S. Sagarika"
    p_g_name.font.name = "Arial"
    p_g_name.font.size = Pt(16)
    p_g_name.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 2: Table of Contents
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide2)
    add_slide_header(slide2, "Table of Contents")

    toc_box = slide2.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.0), Inches(4.5))
    tf_toc = toc_box.text_frame
    tf_toc.word_wrap = True
    
    toc_items = [
        "Abstract",
        "Module description",
        "Software design diagrams",
        "Database design",
        "User interface design",
        "Sample screen shots",
        "Conclusion"
    ]
    for i, item in enumerate(toc_items):
        p = tf_toc.paragraphs[0] if i == 0 else tf_toc.add_paragraph()
        p.text = "➢   " + item
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Abstract & Objectives
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide3)
    add_slide_header(slide3, "Abstract")

    # Abstract Paragraph
    abs_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.3))
    tf_abs = abs_box.text_frame
    tf_abs.word_wrap = True
    tf_abs.margin_left = tf_abs.margin_top = tf_abs.margin_right = tf_abs.margin_bottom = 0
    p_abs = tf_abs.paragraphs[0]
    p_abs.text = "AthleteHub is a web-based platform for managing athletes, coaches, organizations, training activities, tournaments, and performance records."
    p_abs.font.name = "Arial"
    p_abs.font.size = Pt(18)
    p_abs.font.color.rgb = TEXT_DARK
    p_abs.space_after = Pt(8)

    # Objectives Subheader
    obj_lbl = slide3.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(0.5))
    tf_obj_lbl = obj_lbl.text_frame
    tf_obj_lbl.word_wrap = True
    tf_obj_lbl.margin_left = tf_obj_lbl.margin_top = tf_obj_lbl.margin_right = tf_obj_lbl.margin_bottom = 0
    p_obj_lbl = tf_obj_lbl.paragraphs[0]
    p_obj_lbl.text = "Objectives"
    p_obj_lbl.font.name = "Arial"
    p_obj_lbl.font.size = Pt(22)
    p_obj_lbl.font.bold = True
    p_obj_lbl.font.color.rgb = TEXT_RED

    # Objectives Bullet List
    obj_box = slide3.shapes.add_textbox(Inches(1.2), Inches(3.5), Inches(11.3), Inches(3.2))
    tf_obj = obj_box.text_frame
    tf_obj.word_wrap = True
    tf_obj.margin_left = tf_obj.margin_top = tf_obj.margin_right = tf_obj.margin_bottom = 0
    
    objectives = [
        "Centralize athlete and sports information",
        "Manage athletes, coaches, and organizations",
        "Track training and athlete performance",
        "Manage events, tournaments, and matches",
        "Provide secure role-based access"
    ]
    for i, obj in enumerate(objectives):
        p = tf_obj.paragraphs[0] if i == 0 else tf_obj.add_paragraph()
        p.text = "•   " + obj
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: Module Description - Page 1
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide4)
    add_slide_header(slide4, "Module description")

    mod_box1 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf_m1 = mod_box1.text_frame
    tf_m1.word_wrap = True
    tf_m1.margin_left = tf_m1.margin_top = tf_m1.margin_right = tf_m1.margin_bottom = 0

    # Athlete & Organization Management
    p_m1_h = tf_m1.paragraphs[0]
    p_m1_h.text = "Athlete & Organization Management"
    p_m1_h.font.name = "Arial"
    p_m1_h.font.size = Pt(22)
    p_m1_h.font.bold = True
    p_m1_h.font.color.rgb = TEXT_RED
    p_m1_h.space_after = Pt(8)

    bullets1_1 = [
        "Athlete profile management",
        "Organization management",
        "Sports and athlete association",
        "Organization-level data isolation"
    ]
    for bullet in bullets1_1:
        p = tf_m1.add_paragraph()
        p.text = "•   " + bullet
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # Coach Management
    p_m1_h2 = tf_m1.add_paragraph()
    p_m1_h2.text = "Coach Management"
    p_m1_h2.font.name = "Arial"
    p_m1_h2.font.size = Pt(22)
    p_m1_h2.font.bold = True
    p_m1_h2.font.color.rgb = TEXT_RED
    p_m1_h2.space_before = Pt(16)
    p_m1_h2.space_after = Pt(8)

    bullets1_2 = [
        "Coach profile management",
        "Coach assignment to athletes",
        "Assignment status tracking"
    ]
    for bullet in bullets1_2:
        p = tf_m1.add_paragraph()
        p.text = "•   " + bullet
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: Module Description - Page 2
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide5)
    add_slide_header(slide5, "Module description")

    mod_box2 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf_m2 = mod_box2.text_frame
    tf_m2.word_wrap = True
    tf_m2.margin_left = tf_m2.margin_top = tf_m2.margin_right = tf_m2.margin_bottom = 0

    # Training Management
    p_m2_h = tf_m2.paragraphs[0]
    p_m2_h.text = "Training Management"
    p_m2_h.font.name = "Arial"
    p_m2_h.font.size = Pt(22)
    p_m2_h.font.bold = True
    p_m2_h.font.color.rgb = TEXT_RED
    p_m2_h.space_after = Pt(8)

    bullets2_1 = [
        "Training plan creation",
        "Training session scheduling",
        "Athlete attendance tracking"
    ]
    for bullet in bullets2_1:
        p = tf_m2.add_paragraph()
        p.text = "•   " + bullet
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # Event & Tournament Management
    p_m2_h2 = tf_m2.add_paragraph()
    p_m2_h2.text = "Event & Tournament Management"
    p_m2_h2.font.name = "Arial"
    p_m2_h2.font.size = Pt(22)
    p_m2_h2.font.bold = True
    p_m2_h2.font.color.rgb = TEXT_RED
    p_m2_h2.space_before = Pt(16)
    p_m2_h2.space_after = Pt(8)

    bullets2_2 = [
        "Event and tournament creation",
        "Athlete registration",
        "Match scheduling and results"
    ]
    for bullet in bullets2_2:
        p = tf_m2.add_paragraph()
        p.text = "•   " + bullet
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: Module Description - Page 3 (Includes Bulk Import Details)
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide6)
    add_slide_header(slide6, "Module description")

    mod_box3 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf_m3 = mod_box3.text_frame
    tf_m3.word_wrap = True
    tf_m3.margin_left = tf_m3.margin_top = tf_m3.margin_right = tf_m3.margin_bottom = 0

    # Performance Management
    p_m3_h = tf_m3.paragraphs[0]
    p_m3_h.text = "Performance Management"
    p_m3_h.font.name = "Arial"
    p_m3_h.font.size = Pt(22)
    p_m3_h.font.bold = True
    p_m3_h.font.color.rgb = TEXT_RED
    p_m3_h.space_after = Pt(8)

    bullets3_1 = [
        "Performance records",
        "Coach evaluations",
        "Personal Best calculation",
        "Ranking history"
    ]
    for bullet in bullets3_1:
        p = tf_m3.add_paragraph()
        p.text = "•   " + bullet
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # Data Management
    p_m3_h2 = tf_m3.add_paragraph()
    p_m3_h2.text = "Data Management"
    p_m3_h2.font.name = "Arial"
    p_m3_h2.font.size = Pt(22)
    p_m3_h2.font.bold = True
    p_m3_h2.font.color.rgb = TEXT_RED
    p_m3_h2.space_before = Pt(16)
    p_m3_h2.space_after = Pt(6)

    p_dm1 = tf_m3.add_paragraph()
    p_dm1.text = "Data Acquisition & Bulk Import"
    p_dm1.font.name = "Arial"
    p_dm1.font.size = Pt(17)
    p_dm1.font.bold = True
    p_dm1.font.color.rgb = TEXT_DARK
    p_dm1.space_after = Pt(4)

    p_dm2 = tf_m3.add_paragraph()
    p_dm2.text = "AthleteHub supports importing athlete information from: CSV, Excel, JSON"
    p_dm2.font.name = "Arial"
    p_dm2.font.size = Pt(15)
    p_dm2.font.color.rgb = TEXT_MUTED
    p_dm2.space_after = Pt(6)

    p_dm3 = tf_m3.add_paragraph()
    p_dm3.text = "Import Process"
    p_dm3.font.name = "Arial"
    p_dm3.font.size = Pt(17)
    p_dm3.font.bold = True
    p_dm3.font.color.rgb = TEXT_DARK
    p_dm3.space_after = Pt(4)

    p_dm4 = tf_m3.add_paragraph()
    p_dm4.text = "Upload → Mapping → Data Cleaning → Preview → Import"
    p_dm4.font.name = "Arial"
    p_dm4.font.size = Pt(15)
    p_dm4.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 7: Software Design Diagrams
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide7)
    add_slide_header(slide7, "Software design diagrams")

    sw_path = "scratch_extracted/software_design.png"
    if os.path.exists(sw_path):
        # Centered visual block
        slide7.shapes.add_picture(sw_path, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.35))
    else:
        err_box = slide7.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5))
        p = err_box.text_frame.paragraphs[0]
        p.text = "[Software Design Diagram Missing]"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(20)

    # -------------------------------------------------------------
    # SLIDE 8: Database Design
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide8)
    add_slide_header(slide8, "Database design")

    db_path = "scratch_extracted/database_design.png"
    if os.path.exists(db_path):
        # Spans the lower slide area
        slide8.shapes.add_picture(db_path, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.35))
    else:
        err_box = slide8.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5))
        p = err_box.text_frame.paragraphs[0]
        p.text = "[Database Design Diagram Missing]"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(20)

    # -------------------------------------------------------------
    # SLIDE 9: User Interface Design
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide9)
    add_slide_header(slide9, "User interface design")

    ui_path = "scratch_extracted/ui_design.png"
    if os.path.exists(ui_path):
        slide9.shapes.add_picture(ui_path, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.35))
    else:
        err_box = slide9.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5))
        p = err_box.text_frame.paragraphs[0]
        p.text = "[User Interface Design Diagram Missing]"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(20)

    # -------------------------------------------------------------
    # SLIDES 10-13: Sample Screen Shots
    # -------------------------------------------------------------
    screenshots = [
        ("screenshot_dashboard.png", "Sample screen shots"),
        ("screenshot_athletes.png", "Sample screen shots"),
        ("screenshot_training.png", "Sample screen shots"),
        ("screenshot_events.png", "Sample screen shots")
    ]

    for filename, title in screenshots:
        slide = prs.slides.add_slide(blank_layout)
        set_white_bg(slide)
        add_slide_header(slide, title)
        
        path = os.path.join("scratch_extracted", filename)
        if os.path.exists(path):
            # Centered image layout
            slide.shapes.add_picture(path, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.35))
        else:
            err_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5))
            p = err_box.text_frame.paragraphs[0]
            p.text = f"[Screenshot: {filename} Missing]"
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(20)

    # -------------------------------------------------------------
    # SLIDE 14: Conclusion
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide14)
    add_slide_header(slide14, "Conclusion")

    con_box = slide14.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    tf_c = con_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

    con_points = [
        "AthleteHub provides a centralized platform for managing athletes, coaches, organizations, training, events, and performance.",
        "The system implements role-based access control and organization-level data isolation for secure data management.",
        "It integrates training, tournament, match, performance, and ranking management into a single platform.",
    ]
    for i, pt in enumerate(con_points):
        p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
        p.text = "•   " + pt
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    # Future Scope paragraph
    p_fs = tf_c.add_paragraph()
    p_fs.space_before = Pt(10)
    
    r1 = p_fs.add_run()
    r1.text = "Future Scope: "
    r1.font.bold = True
    r1.font.color.rgb = TEXT_RED
    r1.font.size = Pt(18)
    
    r2 = p_fs.add_run()
    r2.text = "AI-powered performance analysis, advanced analytics, athlete recruitment, safeguarding, and live sports integration."
    r2.font.color.rgb = TEXT_DARK
    r2.font.size = Pt(18)

    # -------------------------------------------------------------
    # SLIDE 15: Thank You Slide
    # -------------------------------------------------------------
    slide15 = prs.slides.add_slide(blank_layout)
    set_white_bg(slide15)
    add_sage_bar(slide15)

    ty_box = slide15.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5))
    tf_ty = ty_box.text_frame
    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "THANK YOU"
    p_ty.alignment = PP_ALIGN.CENTER
    p_ty.font.name = "Arial"
    p_ty.font.size = Pt(40)
    p_ty.font.bold = True
    p_ty.font.color.rgb = TEXT_DARK

    # Save presentation
    prs.save("AthleteHub_Presentation.pptx")
    print("Successfully generated AthleteHub_Presentation.pptx!")

if __name__ == "__main__":
    create_presentation()
